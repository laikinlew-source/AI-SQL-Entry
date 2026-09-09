from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


HEADER_SHEET_NAME = "PurchaseInvoiceHeader"
DETAIL_SHEET_NAME = "PurchaseInvoiceDetail"

PROVENANCE_COLUMNS = (
    "SourcePdfFilename",
    "SourceRelativePath",
    "OcrTimestamp",
    "ParserVersion",
    "ProcessingStatus",
    "DocumentId",
    "SourceSha256",
)

HEADER_COLUMNS = (
    "ImportKey",
    "Code",
    "DocNo",
    "DocDate",
    "CurrencyCode",
    "SubTotal",
    "TaxAmt",
    "DocAmt",
    *PROVENANCE_COLUMNS,
)

DETAIL_COLUMNS = (
    "ImportKey",
    "Seq",
    "Account",
    "Description",
    "Qty",
    "UOM",
    "UnitPrice",
    "Tax",
    "TaxAmt",
    "Amount",
    "LineItemId",
    *PROVENANCE_COLUMNS,
)

HEADER_REQUIRED = ("ImportKey", "Code", "DocNo", "DocDate", "CurrencyCode")
DETAIL_REQUIRED = (
    "ImportKey",
    "Seq",
    "Account",
    "Description",
    "Qty",
    "UnitPrice",
    "Tax",
    "Amount",
)
HEADER_NUMERIC = ("SubTotal", "TaxAmt", "DocAmt")
DETAIL_NUMERIC = ("Seq", "Qty", "UnitPrice", "TaxAmt", "Amount")


def _error(code: str, message: str, *, field: str | None = None) -> dict[str, Any]:
    return {"code": code, "message": message, "field": field}


def _headers(worksheet: Any) -> tuple[str, ...]:
    return tuple(
        "" if cell.value is None else str(cell.value)
        for cell in worksheet[1]
    )


def _rows(worksheet: Any, columns: tuple[str, ...]) -> list[dict[str, Any]]:
    return [
        dict(zip(columns, values, strict=True))
        for values in worksheet.iter_rows(
            min_row=2, max_col=len(columns), values_only=True
        )
        if any(value not in (None, "") for value in values)
    ]


def _valid_date(value: Any) -> bool:
    if not isinstance(value, str):
        return False
    try:
        return datetime.strptime(value, "%d/%m/%Y").strftime("%d/%m/%Y") == value
    except ValueError:
        return False


def _is_numeric(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def validate_get_file_3_workbooks(
    header_path: Path, detail_path: Path
) -> dict[str, Any]:
    """Validate the fixed, read-only Get File 3 source workbook contract."""
    errors: list[dict[str, Any]] = []
    check_failures = {
        "sheet_names": False,
        "column_names": False,
        "column_order": False,
        "required_fields": False,
        "date_format": False,
        "numeric_format": False,
        "key_consistency": False,
    }
    header_workbook = load_workbook(header_path, read_only=True, data_only=True)
    detail_workbook = load_workbook(detail_path, read_only=True, data_only=True)
    try:
        if header_workbook.sheetnames != [HEADER_SHEET_NAME]:
            check_failures["sheet_names"] = True
            errors.append(
                _error(
                    "TEMPLATE.HEADER_SHEET",
                    f"Header workbook must contain only sheet '{HEADER_SHEET_NAME}'.",
                )
            )
        if detail_workbook.sheetnames != [DETAIL_SHEET_NAME]:
            check_failures["sheet_names"] = True
            errors.append(
                _error(
                    "TEMPLATE.DETAIL_SHEET",
                    f"Detail workbook must contain only sheet '{DETAIL_SHEET_NAME}'.",
                )
            )
        if check_failures["sheet_names"]:
            return {
                "validator_version": "1.0.0",
                "status": "failed",
                "checks": {
                    name: {"passed": not failed}
                    for name, failed in check_failures.items()
                },
                "errors": errors,
            }

        header_sheet = header_workbook[HEADER_SHEET_NAME]
        detail_sheet = detail_workbook[DETAIL_SHEET_NAME]
        actual_header_columns = _headers(header_sheet)
        actual_detail_columns = _headers(detail_sheet)
        if set(actual_header_columns) != set(HEADER_COLUMNS):
            check_failures["column_names"] = True
        if actual_header_columns != HEADER_COLUMNS:
            check_failures["column_order"] = True
            errors.append(
                _error(
                    "TEMPLATE.HEADER_COLUMNS",
                    "Header workbook columns do not match the fixed contract.",
                )
            )
        if set(actual_detail_columns) != set(DETAIL_COLUMNS):
            check_failures["column_names"] = True
        if actual_detail_columns != DETAIL_COLUMNS:
            check_failures["column_order"] = True
            errors.append(
                _error(
                    "TEMPLATE.DETAIL_COLUMNS",
                    "Detail workbook columns do not match the fixed contract.",
                )
            )

        header_rows = _rows(header_sheet, HEADER_COLUMNS)
        detail_rows = _rows(detail_sheet, DETAIL_COLUMNS)
        if len(header_rows) != 1:
            check_failures["required_fields"] = True
            errors.append(
                _error(
                    "TEMPLATE.HEADER_ROW_COUNT",
                    "Header workbook must contain exactly one data row.",
                )
            )
        if not detail_rows:
            check_failures["required_fields"] = True
            errors.append(
                _error(
                    "TEMPLATE.DETAIL_ROW_COUNT",
                    "Detail workbook must contain at least one data row.",
                )
            )

        for row_number, row in enumerate(header_rows, start=2):
            for field in HEADER_REQUIRED:
                if row[field] in (None, ""):
                    check_failures["required_fields"] = True
                    errors.append(
                        _error(
                            "TEMPLATE.REQUIRED_FIELD",
                            f"Header row {row_number} requires {field}.",
                            field=field,
                        )
                    )
            if not _valid_date(row["DocDate"]):
                check_failures["date_format"] = True
                errors.append(
                    _error(
                        "TEMPLATE.DATE_FORMAT",
                        f"Header row {row_number} DocDate must use dd/mm/yyyy.",
                        field="DocDate",
                    )
                )
            for field in HEADER_NUMERIC:
                if row[field] not in (None, "") and not _is_numeric(row[field]):
                    check_failures["numeric_format"] = True
                    errors.append(
                        _error(
                            "TEMPLATE.NUMERIC_FORMAT",
                            f"Header row {row_number} {field} must be numeric.",
                            field=field,
                        )
                    )

        for row_number, row in enumerate(detail_rows, start=2):
            for field in DETAIL_REQUIRED:
                if row[field] in (None, ""):
                    check_failures["required_fields"] = True
                    errors.append(
                        _error(
                            "TEMPLATE.REQUIRED_FIELD",
                            f"Detail row {row_number} requires {field}.",
                            field=field,
                        )
                    )
            for field in DETAIL_NUMERIC:
                if row[field] not in (None, "") and not _is_numeric(row[field]):
                    check_failures["numeric_format"] = True
                    errors.append(
                        _error(
                            "TEMPLATE.NUMERIC_FORMAT",
                            f"Detail row {row_number} {field} must be numeric.",
                            field=field,
                        )
                    )

        header_keys = {row["ImportKey"] for row in header_rows}
        detail_keys = {row["ImportKey"] for row in detail_rows}
        if header_keys != detail_keys or None in header_keys or "" in header_keys:
            check_failures["key_consistency"] = True
            errors.append(
                _error(
                    "TEMPLATE.KEY_MISMATCH",
                    "Every detail ImportKey must match the single header ImportKey.",
                    field="ImportKey",
                )
            )

        return {
            "validator_version": "1.0.0",
            "status": "passed" if not errors else "failed",
            "checks": {
                name: {"passed": not failed}
                for name, failed in check_failures.items()
            },
            "errors": errors,
        }
    finally:
        header_workbook.close()
        detail_workbook.close()
