from __future__ import annotations

import csv
import json
import os
import re
import shutil
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from typing import Any
from uuid import uuid4

from openpyxl import load_workbook


DATASET_KEYS = {
    "suppliers": "suppliers",
    "tax_codes": "tax_codes",
    "gl_accounts": "gl_accounts",
    "currencies": "currencies",
}


def _normalized_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value).strip().casefold()).strip("_")


def normalize_supplier_name(value: str) -> str:
    """Normalize supplier names for duplicate detection and invoice matching."""
    return re.sub(r"[^a-z0-9]", "", value.casefold())


def _csv_rows(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        reader = csv.DictReader(stream)
        if reader.fieldnames is None:
            return []
        reader.fieldnames = [_normalized_header(name) for name in reader.fieldnames]
        return [dict(row) for row in reader]


def _json_rows(path: Path, dataset: str) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    records = payload.get(dataset) if isinstance(payload, dict) else payload
    if not isinstance(records, list) or any(not isinstance(row, dict) for row in records):
        raise ValueError("JSON input must be a record list or matching snapshot object")
    return records


def _excel_rows(path: Path, sheet_name: str | None) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
        values = worksheet.iter_rows(values_only=True)
        headers = next(values, None)
        if headers is None:
            return []
        normalized_headers = [_normalized_header(value) for value in headers]
        return [
            dict(zip(normalized_headers, row, strict=False))
            for row in values
            if any(value not in (None, "") for value in row)
        ]
    finally:
        workbook.close()


def read_import_rows(
    path: Path, dataset: str, *, sheet_name: str | None = None
) -> tuple[str, list[dict[str, Any]]]:
    """Read rows without applying master-data business validation."""
    suffix = path.suffix.casefold()
    if suffix == ".csv":
        return "csv", _csv_rows(path)
    if suffix == ".json":
        return "json", _json_rows(path, dataset)
    if suffix == ".xlsx":
        return "excel", _excel_rows(path, sheet_name)
    raise ValueError("Supported input formats are CSV, XLSX, and JSON")


def _required_text(row: dict[str, Any], field: str) -> str:
    value = row.get(field)
    if value is None or not str(value).strip():
        raise ValueError(f"Required field '{field}' is empty")
    return str(value).strip()


def _boolean(value: Any, *, default: bool) -> bool:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return value
    normalized = str(value).strip().casefold()
    if normalized in {"true", "yes", "1", "active"}:
        return True
    if normalized in {"false", "no", "0", "inactive"}:
        return False
    raise ValueError(f"Invalid boolean value: {value}")


def _list_value(value: Any) -> list[str]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    return [item.strip() for item in str(value).split(";") if item.strip()]


def _normalize_record(dataset: str, row: dict[str, Any]) -> dict[str, Any]:
    code = _required_text(row, "code")
    active = _boolean(row.get("active"), default=True)
    if dataset == "suppliers":
        name = _required_text(row, "name")
        return {
            "code": code,
            "name": name,
            "normalized_name": normalize_supplier_name(name),
            "aliases": _list_value(row.get("aliases")),
            "active": active,
        }
    if dataset == "tax_codes":
        tax_amount = _required_text(row, "tax_amount").casefold()
        if tax_amount not in {"zero", "nonzero"}:
            raise ValueError("tax_amount must be 'zero' or 'nonzero'")
        return {"code": code, "tax_amount": tax_amount, "active": active}
    if dataset == "gl_accounts":
        return {
            "code": code,
            "name": _required_text(row, "name"),
            "keywords": _list_value(row.get("keywords")),
            "default": _boolean(row.get("default"), default=False),
            "active": active,
        }
    if dataset == "currencies":
        currency_code = code.upper()
        if re.fullmatch(r"[A-Z]{3}", currency_code) is None:
            raise ValueError("Currency code must contain exactly three letters")
        return {"code": currency_code, "active": active}
    raise ValueError(f"Unsupported dataset: {dataset}")


def _duplicates(values: list[tuple[str, str]]) -> list[str]:
    first_values: dict[str, str] = {}
    duplicates: list[str] = []
    for normalized, display in values:
        if normalized in first_values and first_values[normalized] not in duplicates:
            duplicates.append(first_values[normalized])
        else:
            first_values[normalized] = display
    return duplicates


def validate_import_rows(
    dataset: str, rows: list[dict[str, Any]]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, list[str]]]:
    """Validate and normalize reader-independent row dictionaries."""
    if dataset not in DATASET_KEYS:
        raise ValueError(f"Unsupported dataset: {dataset}")
    records: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    for row_number, row in enumerate(rows, start=2):
        try:
            records.append(_normalize_record(dataset, row))
        except (TypeError, ValueError) as error:
            errors.append({"row": row_number, "message": str(error)})

    duplicate_report = {"supplier_codes": [], "supplier_names": []}
    if dataset == "suppliers":
        duplicate_report["supplier_codes"] = _duplicates(
            [(record["code"].casefold(), record["code"]) for record in records]
        )
        duplicate_report["supplier_names"] = _duplicates(
            [
                (record["normalized_name"], record["normalized_name"])
                for record in records
            ]
        )
    return records, errors, duplicate_report


def _read_existing_count(snapshot_path: Path, dataset: str) -> int:
    if not snapshot_path.exists():
        return 0
    payload = json.loads(snapshot_path.read_text(encoding="utf-8"))
    records = payload.get(dataset, [])
    if not isinstance(records, list):
        raise ValueError(f"Existing {snapshot_path.name} has an invalid record list")
    return len(records)


def _timestamp_text(timestamp: datetime) -> tuple[str, str]:
    utc_timestamp = timestamp.astimezone(timezone.utc)
    return (
        utc_timestamp.isoformat().replace("+00:00", "Z"),
        utc_timestamp.strftime("%Y%m%dT%H%M%S%fZ"),
    )


def _write_snapshot(path: Path, dataset: str, records: list[dict[str, Any]]) -> None:
    temporary_path = path.parent / f"{path.name}.{uuid4().hex}.tmp"
    payload = {"schema_version": "1.0.0", dataset: records}
    with temporary_path.open("w", encoding="utf-8", newline="\n") as stream:
        json.dump(payload, stream, indent=2, ensure_ascii=False)
        stream.write("\n")
        stream.flush()
        os.fsync(stream.fileno())
    os.replace(temporary_path, path)


def build_snapshot(
    dataset: str,
    input_path: Path,
    config_dir: Path,
    *,
    dry_run: bool = False,
    minimum_retention_ratio: Decimal = Decimal("0.50"),
    sheet_name: str | None = None,
    timestamp: datetime | None = None,
) -> dict[str, Any]:
    """Validate an import and safely replace one local master-data snapshot."""
    if minimum_retention_ratio < 0 or minimum_retention_ratio > 1:
        raise ValueError("minimum_retention_ratio must be between 0 and 1")
    timestamp = timestamp or datetime.now(timezone.utc)
    timestamp_value, timestamp_token = _timestamp_text(timestamp)
    snapshot_path = config_dir / f"{dataset}.json"
    existing_count = _read_existing_count(snapshot_path, dataset)
    try:
        source_format, rows = read_import_rows(
            input_path, dataset, sheet_name=sheet_name
        )
    except Exception as error:
        source_format = {
            ".csv": "csv",
            ".json": "json",
            ".xlsx": "excel",
        }.get(input_path.suffix.casefold(), "unsupported")
        return {
            "report_version": "1.0.0",
            "dataset": dataset,
            "input_path": str(input_path),
            "source_format": source_format,
            "snapshot_path": str(snapshot_path),
            "generated_at": timestamp_value,
            "dry_run": dry_run,
            "minimum_retention_ratio": format(minimum_retention_ratio, "f"),
            "statistics": {
                "rows_read": 0,
                "rows_accepted": 0,
                "rows_rejected": 0,
                "existing_rows": existing_count,
                "retention_ratio": "0.0000" if existing_count else "1.0000",
            },
            "duplicates": {"supplier_codes": [], "supplier_names": []},
            "errors": [{"row": None, "message": str(error)}],
            "warnings": [],
            "would_replace_snapshot": False,
            "snapshot_written": False,
            "backup_path": None,
            "sql_account_access": "none",
            "status": "invalid",
        }
    records, errors, duplicates = validate_import_rows(dataset, rows)
    retention_ratio = (
        Decimal(len(records)) / Decimal(existing_count)
        if existing_count
        else Decimal("1")
    )
    duplicate_count = sum(len(values) for values in duplicates.values())
    invalid = bool(errors or duplicate_count)
    blocked = (
        not invalid
        and existing_count > 0
        and retention_ratio < minimum_retention_ratio
    )
    warnings: list[str] = []
    if blocked:
        warnings.append(
            "Imported row retention is below the configured minimum; snapshot replacement stopped."
        )

    report: dict[str, Any] = {
        "report_version": "1.0.0",
        "dataset": dataset,
        "input_path": str(input_path),
        "source_format": source_format,
        "snapshot_path": str(snapshot_path),
        "generated_at": timestamp_value,
        "dry_run": dry_run,
        "minimum_retention_ratio": format(minimum_retention_ratio, "f"),
        "statistics": {
            "rows_read": len(rows),
            "rows_accepted": len(records),
            "rows_rejected": len(errors),
            "existing_rows": existing_count,
            "retention_ratio": f"{retention_ratio:.4f}",
        },
        "duplicates": duplicates,
        "errors": errors,
        "warnings": warnings,
        "would_replace_snapshot": not invalid and not blocked,
        "snapshot_written": False,
        "backup_path": None,
        "sql_account_access": "none",
        "status": "invalid" if invalid else "blocked" if blocked else "valid",
    }
    if invalid or blocked or dry_run:
        return report

    config_dir.mkdir(parents=True, exist_ok=True)
    if snapshot_path.exists():
        backup_dir = config_dir / "backups"
        backup_dir.mkdir(parents=True, exist_ok=True)
        backup_path = backup_dir / f"{dataset}-{timestamp_token}.json"
        shutil.copy2(snapshot_path, backup_path)
        report["backup_path"] = str(backup_path)
    _write_snapshot(snapshot_path, dataset, records)
    report["snapshot_written"] = True
    report["status"] = "written"
    return report
