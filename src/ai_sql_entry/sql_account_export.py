from __future__ import annotations

import json
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from statistics import fmean
from typing import Any

from openpyxl import Workbook

from .sql_account_template import (
    DETAIL_COLUMNS,
    DETAIL_SHEET_NAME,
    HEADER_COLUMNS,
    HEADER_SHEET_NAME,
    validate_get_file_3_workbooks,
)
from .sql_account_validation import SqlAccountMasterData, validate_invoice


MONEY_FORMAT = "0.00########"
NUMBER_FORMAT = "0.################"


def _effective_value(field: dict[str, Any]) -> Any:
    reviewed = field.get("reviewed", {})
    if reviewed.get("status") in {"corrected", "supplied"}:
        return reviewed.get("value")
    return field["extracted"]["normalized_value"]


def _decimal(field: dict[str, Any]) -> Decimal:
    value = _effective_value(field)
    if isinstance(value, dict):
        value = value["value"]
    return Decimal(str(value))


def _provenance(canonical: dict[str, Any]) -> list[Any]:
    return [
        canonical["source"]["original_filename"],
        canonical["source"]["source_relative_path"],
        canonical["processing"]["ocr_timestamp"],
        canonical["processing"]["parser_version"],
        canonical["processing"]["processing_status"],
        canonical["document_id"],
        canonical["source"]["sha256"],
    ]


def _missing_mappings(mapping_report: dict[str, Any]) -> list[dict[str, Any]]:
    checks = mapping_report["checks"]
    missing: list[dict[str, Any]] = []
    for name in ("supplier", "tax_code", "currency"):
        check = checks[name]
        if check["status"] != "found":
            missing.append(
                {
                    "mapping": name,
                    "field_path": check["field_path"],
                    "source_value": check["source_value"],
                    "reason": check["reason"],
                }
            )
    for line in checks["gl_account"]["line_items"]:
        if line["status"] != "found":
            missing.append(
                {
                    "mapping": "gl_account",
                    "field_path": line["field_path"],
                    "line_item_id": line["line_item_id"],
                    "source_value": line["source_value"],
                    "reason": line["reason"],
                }
            )
    return missing


def _arithmetic(canonical: dict[str, Any]) -> dict[str, Any]:
    tolerance = Decimal("0.01")
    subtotal = _decimal(canonical["financial_summary"]["subtotal"])
    tax = _decimal(canonical["financial_summary"]["tax"])
    grand_total = _decimal(canonical["financial_summary"]["grand_total"])
    total_difference = grand_total - (subtotal + tax)
    totals = {
        "passed": abs(total_difference) <= tolerance,
        "expected": format(subtotal + tax, "f"),
        "actual": format(grand_total, "f"),
        "difference": format(total_difference, "f"),
        "tolerance": "0.01",
    }
    sst_difference = tax - (grand_total - subtotal)
    sst = {
        "passed": tax >= 0 and abs(sst_difference) <= tolerance,
        "expected": format(grand_total - subtotal, "f"),
        "actual": format(tax, "f"),
        "difference": format(sst_difference, "f"),
        "tolerance": "0.01",
    }

    line_results: list[dict[str, Any]] = []
    line_sum = Decimal("0")
    for index, item in enumerate(canonical["line_items"]):
        quantity = _decimal(item["quantity"])
        unit_price = _decimal(item["unit_price"])
        actual = _decimal(item["line_total"])
        expected = quantity * unit_price
        difference = actual - expected
        line_sum += actual
        line_results.append(
            {
                "line_item_id": item.get("line_item_id"),
                "field_path": f"/line_items/{index}/line_total",
                "passed": abs(difference) <= tolerance,
                "expected": format(expected, "f"),
                "actual": format(actual, "f"),
                "difference": format(difference, "f"),
            }
        )
    subtotal_difference = line_sum - subtotal
    failures = [line for line in line_results if not line["passed"]]
    if abs(subtotal_difference) > tolerance:
        failures.append(
            {
                "line_item_id": None,
                "field_path": "/financial_summary/subtotal",
                "passed": False,
                "expected": format(line_sum, "f"),
                "actual": format(subtotal, "f"),
                "difference": format(subtotal_difference, "f"),
            }
        )
    line_totals = {
        "passed": bool(line_results)
        and all(line["passed"] for line in line_results)
        and abs(subtotal_difference) <= tolerance,
        "line_results": line_results,
        "line_sum": format(line_sum, "f"),
        "document_subtotal": format(subtotal, "f"),
        "subtotal_difference": format(subtotal_difference, "f"),
        "tolerance": "0.01",
        "failures": failures,
    }
    return {"totals": totals, "sst": sst, "line_totals": line_totals}


def _confidence_fields(canonical: dict[str, Any]) -> list[dict[str, Any]]:
    fields = [
        canonical["parties"]["issuer"]["name"],
        canonical["document_details"]["document_number"],
        canonical["document_details"]["issue_date"],
        canonical["document_details"]["currency"],
        canonical["financial_summary"]["subtotal"],
        canonical["financial_summary"]["tax"],
        canonical["financial_summary"]["grand_total"],
    ]
    for item in canonical["line_items"]:
        fields.extend(
            [item["description"], item["quantity"], item["unit_price"], item["line_total"]]
        )
    return fields


def _ocr_score(canonical: dict[str, Any]) -> dict[str, Any]:
    fields = _confidence_fields(canonical)
    confidences = [
        field["extracted"].get("confidence")
        for field in fields
        if field.get("extracted") is not None
        and field["extracted"].get("confidence") is not None
    ]
    complete = len(confidences) == len(fields) and bool(fields)
    score = round(fmean(float(value) for value in confidences) * 100, 2) if complete else None
    return {
        "score": score,
        "threshold": 80.0,
        "status": "passed" if score is not None and score >= 80.0 else "failed",
        "fields_scored": len(confidences),
        "fields_required": len(fields),
        "reason": None
        if complete
        else "One or more required extracted fields have no OCR confidence value.",
    }


def _write_workbook(
    path: Path,
    sheet_name: str,
    columns: tuple[str, ...],
    rows: list[list[Any]],
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(columns)
    for row in rows:
        worksheet.append(row)
    for column_name in ("SubTotal", "TaxAmt", "DocAmt", "Qty", "UnitPrice", "Amount"):
        if column_name not in columns:
            continue
        column_index = columns.index(column_name) + 1
        for row_index in range(2, worksheet.max_row + 1):
            worksheet.cell(row_index, column_index).number_format = (
                MONEY_FORMAT if column_name != "Qty" else NUMBER_FORMAT
            )
    workbook.save(path)
    workbook.close()


def _write_json(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )


def _validation_report(
    canonical: dict[str, Any],
    mapping_report: dict[str, Any],
    arithmetic: dict[str, Any],
    template_report: dict[str, Any],
    *,
    generated_at: str,
) -> dict[str, Any]:
    mapping_checks = mapping_report["checks"]
    ocr = _ocr_score(canonical)
    scores = {
        "OCRAccuracy": ocr,
        "SupplierMatch": {
            "score": 100.0 if mapping_checks["supplier"]["status"] == "found" else 0.0
        },
        "TaxMatch": {
            "score": 100.0 if mapping_checks["tax_code"]["status"] == "found" else 0.0
        },
        "GLMatch": {
            "score": 100.0 if mapping_checks["gl_account"]["status"] == "found" else 0.0
        },
        "CurrencyMatch": {
            "score": 100.0 if mapping_checks["currency"]["status"] == "found" else 0.0
        },
        "TemplateValidation": {
            "score": 100.0 if template_report["status"] == "passed" else 0.0
        },
    }
    component_scores = [component["score"] or 0.0 for component in scores.values()]
    arithmetic_passed = all(result["passed"] for result in arithmetic.values())
    mappings_passed = not mapping_report["missing_mappings"]
    review_passed = (
        canonical.get("review", {}).get("status") == "approved"
        and canonical.get("review", {}).get("decision") == "approved"
    )
    ready = (
        ocr["status"] == "passed"
        and mappings_passed
        and arithmetic_passed
        and review_passed
        and template_report["status"] == "passed"
    )
    scores["OverallScore"] = round(fmean(component_scores), 2)
    scores["OverallReadyForImport"] = "YES" if ready else "NO"
    provenance = {
        "source_pdf_filename": canonical["source"]["original_filename"],
        "source_relative_path": canonical["source"]["source_relative_path"],
        "ocr_timestamp": canonical["processing"]["ocr_timestamp"],
        "parser_version": canonical["processing"]["parser_version"],
        "processing_status": canonical["processing"]["processing_status"],
        "document_id": canonical["document_id"],
        "source_sha256": canonical["source"]["sha256"],
    }
    return {
        "report_version": "1.0.0",
        **provenance,
        "generated_at": generated_at,
        "target": "SQLAccXLSnMDBImp Get File 3",
        "access_mode": "file_generation_only",
        "sql_account_database_accessed": False,
        "sql_account_ui_automated": False,
        "validations": {
            **arithmetic,
            "supplier_code": {
                "passed": mapping_checks["supplier"]["status"] == "found"
            },
            "tax_code": {
                "passed": mapping_checks["tax_code"]["status"] == "found"
            },
            "gl_account": {
                "passed": mapping_checks["gl_account"]["status"] == "found"
            },
            "currency": {
                "passed": mapping_checks["currency"]["status"] == "found"
            },
            "human_review": {"passed": review_passed},
            "template_validation": template_report,
        },
        "missing_mappings": mapping_report["missing_mappings"],
        "import_readiness_score": scores,
        "ReadyForImport": "YES" if ready else "NO",
    }


def _write_markdown_reports(output_dir: Path, report: dict[str, Any]) -> None:
    ready = report["ReadyForImport"]
    score = report["import_readiness_score"]["OverallScore"]
    missing = report["missing_mappings"]
    summary_lines = [
        "# SQL Account Import Summary",
        "",
        f"- Target: SQLAccXLSnMDBImp Get File 3",
        f"- Document ID: `{report['document_id']}`",
        f"- Source PDF: `{report['source_pdf_filename']}`",
        f"- Source relative path: `{report['source_relative_path']}`",
        f"- OCR timestamp: `{report['ocr_timestamp']}`",
        f"- Parser version: `{report['parser_version']}`",
        f"- Processing status: `{report['processing_status']}`",
        f"- Overall score: {score}",
        f"- ReadyForImport: **{ready}**",
        "",
    ]
    if missing:
        summary_lines.extend(["## Missing mappings", ""])
        summary_lines.extend(
            f"- `{item['field_path']}`: {item['reason']}" for item in missing
        )
        summary_lines.append("")
    (output_dir / "Import_Summary.md").write_text(
        "\n".join(summary_lines), encoding="utf-8"
    )

    checklist = f"""# SQL Account Import Checklist

- Source PDF: `{report['source_pdf_filename']}`
- Source relative path: `{report['source_relative_path']}`
- OCR timestamp: `{report['ocr_timestamp']}`
- Parser version: `{report['parser_version']}`
- Processing status: `{report['processing_status']}`
- Document ID: `{report['document_id']}`

- [ ] Confirm `Validation_Report.json` shows `ReadyForImport` as `YES`.
- [ ] Confirm the source PDF, supplier, invoice number, date, currency, totals, and line items.
- [ ] Open the official SQLAccXLSnMDBImp utility manually.
- [ ] From the Get File menu, select `Get File 3...`.
- [ ] Select `Purchase Invoice Header.xlsx` and sheet `PurchaseInvoiceHeader`.
- [ ] Select `Purchase Invoice Detail.xlsx` and sheet `PurchaseInvoiceDetail`.
- [ ] Select `ImportKey` as the common Key Field for both files.
- [ ] Map Header fields: `Code`, `DocNo`, `DocDate`, `CurrencyCode`, `SubTotal`, `TaxAmt`, `DocAmt`.
- [ ] Map Detail fields: `Seq`, `Account`, `Description`, `Qty`, `UOM` when present, `UnitPrice`, `Tax`, `Amount`.
- [ ] Leave provenance fields and `LineItemId` unmapped.
- [ ] Click `Get Data`, then click `Verify` and resolve every reported issue.
- [ ] Do not click `Post To A/c` without separate human authorization.

AI-SQL-Entry generates files only. It does not access SQL Account, automate its UI, or post accounting data.
"""
    (output_dir / "Import_Checklist.md").write_text(checklist, encoding="utf-8")


def write_import_package(
    canonical: dict[str, Any],
    master_data: SqlAccountMasterData,
    output_dir: Path,
    *,
    generated_at: str,
) -> dict[str, Any]:
    """Generate a fail-closed Get File 3 package without SQL Account access."""
    output_dir.mkdir(parents=True, exist_ok=True)
    mapping_report = validate_invoice(
        canonical, master_data, validated_at=generated_at
    )
    mapping_report["target"] = "SQLAccXLSnMDBImp Get File 3"
    mapping_report["missing_mappings"] = _missing_mappings(mapping_report)
    mapping_report["all_mappings_found"] = not mapping_report["missing_mappings"]
    _write_json(output_dir / "Mapping_Report.json", mapping_report)

    arithmetic = _arithmetic(canonical)
    ocr = _ocr_score(canonical)
    review_passed = (
        canonical.get("review", {}).get("status") == "approved"
        and canonical.get("review", {}).get("decision") == "approved"
    )
    preflight_passed = (
        mapping_report["all_mappings_found"]
        and all(result["passed"] for result in arithmetic.values())
        and ocr["status"] == "passed"
        and review_passed
    )
    template_report: dict[str, Any] = {
        "validator_version": "1.0.0",
        "status": "not_run",
        "checks": {},
        "errors": [],
        "reason": "Preflight validation did not pass; import workbooks were not generated.",
    }

    if preflight_passed:
        details = canonical["document_details"]
        summary = canonical["financial_summary"]
        checks = mapping_report["checks"]
        key = canonical["document_id"]
        document_date = datetime.strptime(
            str(_effective_value(details["issue_date"])), "%Y-%m-%d"
        ).strftime("%d/%m/%Y")
        header_row = [
            key,
            checks["supplier"]["mapped_code"],
            _effective_value(details["document_number"]),
            document_date,
            checks["currency"]["mapped_code"],
            _decimal(summary["subtotal"]),
            _decimal(summary["tax"]),
            _decimal(summary["grand_total"]),
            *_provenance(canonical),
        ]
        tax_code = checks["tax_code"]["mapped_code"]
        detail_rows = []
        for position, item in enumerate(canonical["line_items"], start=1):
            gl_mapping = checks["gl_account"]["line_items"][position - 1]
            unit_of_measure = item.get("unit_of_measure")
            detail_rows.append(
                [
                    key,
                    position,
                    gl_mapping["mapped_code"],
                    _effective_value(item["description"]),
                    _decimal(item["quantity"]),
                    _effective_value(unit_of_measure) if unit_of_measure else "",
                    _decimal(item["unit_price"]),
                    tax_code,
                    "",
                    _decimal(item["line_total"]),
                    item["line_item_id"],
                    *_provenance(canonical),
                ]
            )
        header_path = output_dir / "Purchase Invoice Header.xlsx"
        detail_path = output_dir / "Purchase Invoice Detail.xlsx"
        _write_workbook(
            header_path, HEADER_SHEET_NAME, HEADER_COLUMNS, [header_row]
        )
        _write_workbook(
            detail_path, DETAIL_SHEET_NAME, DETAIL_COLUMNS, detail_rows
        )
        template_report = validate_get_file_3_workbooks(header_path, detail_path)

    validation_report = _validation_report(
        canonical,
        mapping_report,
        arithmetic,
        template_report,
        generated_at=generated_at,
    )
    _write_json(output_dir / "Validation_Report.json", validation_report)
    _write_markdown_reports(output_dir, validation_report)
    return validation_report
