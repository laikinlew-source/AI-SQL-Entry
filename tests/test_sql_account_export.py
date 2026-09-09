from __future__ import annotations

import json
import shutil
import unittest
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

from ai_sql_entry.canonical import build_canonical
from ai_sql_entry.extraction import extract_invoice
from ai_sql_entry.sql_account_export import write_import_package
from ai_sql_entry.sql_account_validation import (
    CurrencyRecord,
    GlAccountRecord,
    SqlAccountMasterData,
    SupplierRecord,
    TaxCodeRecord,
)
from ai_sql_entry.sql_account_template import (
    DETAIL_COLUMNS,
    DETAIL_SHEET_NAME,
    HEADER_COLUMNS,
    HEADER_SHEET_NAME,
)


FIXTURES = Path(__file__).parent / "fixtures"
RUNTIME = Path(__file__).parent / "runtime" / "sql-import-package"


def _canonical_invoice() -> dict[str, object]:
    invoice = extract_invoice(
        (FIXTURES / "invoice_ocr.txt").read_text(encoding="utf-8")
    )
    canonical = build_canonical(
        invoice,
        document_id="doc-test-0001",
        source_filename="invoice.pdf",
        source_relative_path="invoices/invoice.pdf",
        source_sha256="a" * 64,
        page_count=1,
        created_at=datetime(2026, 7, 21, 4, 15, 30, tzinfo=timezone.utc),
        ocr_timestamp="2026-07-21T04:15:31Z",
        parser_version="0.3.0",
        processing_status="extraction_succeeded",
    )
    canonical["review"]["status"] = "approved"
    canonical["review"]["decision"] = "approved"
    required_fields = [
        canonical["parties"]["issuer"]["name"],
        canonical["document_details"]["document_number"],
        canonical["document_details"]["issue_date"],
        canonical["document_details"]["currency"],
        canonical["financial_summary"]["subtotal"],
        canonical["financial_summary"]["tax"],
        canonical["financial_summary"]["grand_total"],
    ]
    for item in canonical["line_items"]:
        required_fields.extend(
            [item["description"], item["quantity"], item["unit_price"], item["line_total"]]
        )
    for field in required_fields:
        field["extracted"]["confidence"] = 0.95
    return canonical


def _matching_master_data() -> SqlAccountMasterData:
    return SqlAccountMasterData(
        suppliers=(
            SupplierRecord(code="SUP-ACORN", name="Acorn Office Supplies Sdn Bhd"),
        ),
        tax_codes=(TaxCodeRecord(code="SST-6", tax_amount="nonzero"),),
        gl_accounts=(
            GlAccountRecord(
                code="5000-PURCHASES",
                name="Purchases",
                keywords=("desk",),
            ),
        ),
        currencies=(CurrencyRecord(code="MYR"),),
    )


class SqlAccountImportPackageTests(unittest.TestCase):
    def setUp(self) -> None:
        shutil.rmtree(RUNTIME, ignore_errors=True)

    def tearDown(self) -> None:
        shutil.rmtree(RUNTIME, ignore_errors=True)

    def test_writes_two_valid_get_file_3_workbooks_and_four_reports(self) -> None:
        result = write_import_package(
            _canonical_invoice(),
            _matching_master_data(),
            RUNTIME,
            generated_at="2026-07-21T05:00:00Z",
        )

        self.assertEqual(
            {path.name for path in RUNTIME.iterdir()},
            {
                "Purchase Invoice Header.xlsx",
                "Purchase Invoice Detail.xlsx",
                "Mapping_Report.json",
                "Validation_Report.json",
                "Import_Checklist.md",
                "Import_Summary.md",
            },
        )
        self.assertEqual(result["ReadyForImport"], "YES")

        header = load_workbook(
            RUNTIME / "Purchase Invoice Header.xlsx", read_only=True, data_only=True
        )
        detail = load_workbook(
            RUNTIME / "Purchase Invoice Detail.xlsx", read_only=True, data_only=True
        )
        try:
            header_sheet = header[HEADER_SHEET_NAME]
            detail_sheet = detail[DETAIL_SHEET_NAME]
            self.assertEqual(tuple(cell.value for cell in header_sheet[1]), HEADER_COLUMNS)
            self.assertEqual(tuple(cell.value for cell in detail_sheet[1]), DETAIL_COLUMNS)
            header_row = tuple(cell.value for cell in header_sheet[2])
            detail_row = tuple(cell.value for cell in detail_sheet[2])
        finally:
            header.close()
            detail.close()

        self.assertEqual(header_row[0], "doc-test-0001")
        self.assertEqual(detail_row[0], header_row[0])
        self.assertEqual(
            header_row[1:5],
            ("SUP-ACORN", "INV-0042", "18/07/2026", "MYR"),
        )
        self.assertEqual(detail_row[2], "5000-PURCHASES")
        self.assertEqual(detail_row[7], "SST-6")
        self.assertIsInstance(header_row[5], (int, float))
        self.assertIsInstance(detail_row[4], (int, float))

        mapping = json.loads((RUNTIME / "Mapping_Report.json").read_text(encoding="utf-8"))
        validation = json.loads(
            (RUNTIME / "Validation_Report.json").read_text(encoding="utf-8")
        )
        self.assertEqual(mapping["checks"]["supplier"]["mapped_code"], "SUP-ACORN")
        self.assertEqual(mapping["missing_mappings"], [])
        self.assertEqual(validation["ReadyForImport"], "YES")
        self.assertEqual(
            validation["import_readiness_score"]["OCRAccuracy"]["score"], 95.0
        )
        self.assertEqual(
            validation["import_readiness_score"]["TemplateValidation"]["score"],
            100.0,
        )
        self.assertEqual(
            validation["import_readiness_score"]["OverallReadyForImport"], "YES"
        )
        self.assertTrue(validation["validations"]["totals"]["passed"])
        self.assertTrue(validation["validations"]["sst"]["passed"])
        self.assertTrue(validation["validations"]["line_totals"]["passed"])
        self.assertEqual(validation["source_pdf_filename"], "invoice.pdf")
        for report_name in ("Import_Checklist.md", "Import_Summary.md"):
            markdown = (RUNTIME / report_name).read_text(encoding="utf-8")
            self.assertIn("invoice.pdf", markdown)
            self.assertIn("invoices/invoice.pdf", markdown)
            self.assertIn("2026-07-21T04:15:31Z", markdown)
            self.assertIn("0.3.0", markdown)
            self.assertIn("extraction_succeeded", markdown)

    def test_mapping_failures_generate_reports_but_no_excel(self) -> None:
        result = write_import_package(
            _canonical_invoice(),
            SqlAccountMasterData(),
            RUNTIME,
            generated_at="2026-07-21T05:00:00Z",
        )

        self.assertEqual(result["ReadyForImport"], "NO")
        self.assertFalse((RUNTIME / "Purchase Invoice Header.xlsx").exists())
        self.assertFalse((RUNTIME / "Purchase Invoice Detail.xlsx").exists())
        self.assertTrue((RUNTIME / "Mapping_Report.json").exists())
        self.assertTrue((RUNTIME / "Validation_Report.json").exists())
        mapping = json.loads((RUNTIME / "Mapping_Report.json").read_text(encoding="utf-8"))
        self.assertEqual(
            {item["mapping"] for item in mapping["missing_mappings"]},
            {"supplier", "tax_code", "gl_account", "currency"},
        )
        validation = json.loads(
            (RUNTIME / "Validation_Report.json").read_text(encoding="utf-8")
        )
        self.assertEqual(validation["ReadyForImport"], "NO")
        self.assertEqual(
            validation["validations"]["template_validation"]["status"], "not_run"
        )

    def test_arithmetic_failure_reports_exact_line_field_and_blocks_workbooks(self) -> None:
        canonical = _canonical_invoice()
        canonical["line_items"][0]["line_total"]["extracted"]["normalized_value"][
            "value"
        ] = "99.00"

        result = write_import_package(
            canonical,
            _matching_master_data(),
            RUNTIME,
            generated_at="2026-07-21T05:00:00Z",
        )

        self.assertEqual(result["ReadyForImport"], "NO")
        self.assertFalse((RUNTIME / "Purchase Invoice Header.xlsx").exists())
        failures = result["validations"]["line_totals"]["failures"]
        self.assertEqual(failures[0]["field_path"], "/line_items/0/line_total")
        self.assertIn("expected", failures[0])
        self.assertIn("actual", failures[0])


if __name__ == "__main__":
    unittest.main()
