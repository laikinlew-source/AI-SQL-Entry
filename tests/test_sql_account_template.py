from __future__ import annotations

import shutil
import unittest
from pathlib import Path

from openpyxl import Workbook

from ai_sql_entry.sql_account_template import (
    DETAIL_COLUMNS,
    DETAIL_SHEET_NAME,
    HEADER_COLUMNS,
    HEADER_SHEET_NAME,
    validate_get_file_3_workbooks,
)


RUNTIME = Path(__file__).parent / "runtime" / "template-validator"


def _write_workbook(path: Path, sheet_name: str, columns: tuple[str, ...], row: list[object]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.append(columns)
    worksheet.append(row)
    workbook.save(path)
    workbook.close()


class SqlAccountTemplateValidatorTests(unittest.TestCase):
    def setUp(self) -> None:
        shutil.rmtree(RUNTIME, ignore_errors=True)
        RUNTIME.mkdir(parents=True)
        self.header_path = RUNTIME / "Purchase Invoice Header.xlsx"
        self.detail_path = RUNTIME / "Purchase Invoice Detail.xlsx"

    def tearDown(self) -> None:
        shutil.rmtree(RUNTIME, ignore_errors=True)

    def _write_valid_pair(self) -> None:
        _write_workbook(
            self.header_path,
            HEADER_SHEET_NAME,
            HEADER_COLUMNS,
            [
                "doc-abc",
                "SUP-001",
                "INV-0042",
                "18/07/2026",
                "MYR",
                100.00,
                6.00,
                106.00,
                "invoice.pdf",
                "invoices/invoice.pdf",
                "2026-07-21T04:15:31Z",
                "0.3.0",
                "extraction_succeeded",
                "doc-abc",
                "a" * 64,
            ],
        )
        _write_workbook(
            self.detail_path,
            DETAIL_SHEET_NAME,
            DETAIL_COLUMNS,
            [
                "doc-abc",
                1,
                "5000-PURCHASES",
                "Desk Set",
                2,
                "",
                50.00,
                "SST-6",
                "",
                100.00,
                "line-0001",
                "invoice.pdf",
                "invoices/invoice.pdf",
                "2026-07-21T04:15:31Z",
                "0.3.0",
                "extraction_succeeded",
                "doc-abc",
                "a" * 64,
            ],
        )

    def test_accepts_the_fixed_get_file_3_contract(self) -> None:
        self._write_valid_pair()

        report = validate_get_file_3_workbooks(self.header_path, self.detail_path)

        self.assertEqual(report["status"], "passed")
        self.assertEqual(report["errors"], [])
        self.assertTrue(all(check["passed"] for check in report["checks"].values()))

    def test_reports_column_date_numeric_and_key_failures(self) -> None:
        self._write_valid_pair()
        from openpyxl import load_workbook

        header = load_workbook(self.header_path)
        header_sheet = header[HEADER_SHEET_NAME]
        header_sheet.cell(1, 2).value = "WrongCodeColumn"
        header_sheet.cell(2, 4).value = "2026-07-18"
        header_sheet.cell(2, 6).value = "RM 100.00"
        header.save(self.header_path)
        header.close()

        detail = load_workbook(self.detail_path)
        detail[DETAIL_SHEET_NAME].cell(2, 1).value = "different-key"
        detail.save(self.detail_path)
        detail.close()

        report = validate_get_file_3_workbooks(self.header_path, self.detail_path)

        self.assertEqual(report["status"], "failed")
        codes = {error["code"] for error in report["errors"]}
        self.assertIn("TEMPLATE.HEADER_COLUMNS", codes)
        self.assertIn("TEMPLATE.DATE_FORMAT", codes)
        self.assertIn("TEMPLATE.NUMERIC_FORMAT", codes)
        self.assertIn("TEMPLATE.KEY_MISMATCH", codes)


if __name__ == "__main__":
    unittest.main()
